"""Tests for helpers.run_sheet — the workbook engine.

The bar these hold: the student/assessor split (nothing assessor-only leaks into a student copy),
and the scaffolding dials on a capture table (`given`, `exemplar`, `blank_rows`), because those are
what decide how much of an answer the instrument hands over.
"""
import pathlib
import sys

sys.path.insert(0, str(next(
    d for d in pathlib.Path(__file__).resolve().parents
    if (d / "helpers" / "__init__.py").exists()
)))

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from helpers.run_sheet import (  # noqa: E402
    design_table, element, response_slot, screenshot_slot, standard_line, uoc_line, MODEL, UOC,
)


def _cells(doc, table=0):
    return [[c.text for c in row.cells] for row in doc.tables[table].rows]


# ---------------------------------------------------------------- the student/assessor split

def test_uoc_line_assessor_only():
    doc = Document()
    uoc_line(doc, ["ICTCLD502 PC 2.2"], "student")
    assert doc.paragraphs == []                   # nothing rendered at all
    uoc_line(doc, ["ICTCLD502 PC 2.2"], "assessor")
    assert doc.paragraphs[-1].text == "Evidences: [ICTCLD502 PC 2.2]"
    assert str(doc.paragraphs[-1].runs[0].font.color.rgb) == UOC


def test_standard_line_assessor_only():
    doc = Document()
    standard_line(doc, "the student names three single points of failure", "student")
    assert all("Satisfactory when" not in p.text for p in doc.paragraphs)
    standard_line(doc, "the student names three single points of failure", "assessor")
    assert doc.paragraphs[-1].text.startswith("Satisfactory when ")


def test_element_student_copy_carries_no_marking_apparatus():
    """The whole point of one definition rendered twice: the student copy cannot leak the answer."""
    doc = Document()
    el = dict(n=3, title="Single points of failure", prompt="Identify every one.",
              uoc=["ICTCLD502 PC 2.2"], standard="all three are named",
              table=(["Component", "Consequence"], [["the single NAT gateway", "no outbound"]]))
    element(doc, lambda t: doc.add_paragraph(t), el, "student")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ICTCLD502" not in text and "Satisfactory when" not in text
    assert "the single NAT gateway" not in [c for row in _cells(doc) for c in row]


def test_element_assessor_copy_carries_the_model_answer():
    doc = Document()
    el = dict(n=3, title="Single points of failure", prompt="Identify every one.",
              uoc=["ICTCLD502 PC 2.2"], standard="all three are named",
              table=(["Component", "Consequence"], [["the single NAT gateway", "no outbound"]]))
    element(doc, lambda t: doc.add_paragraph(t), el, "assessor")
    assert _cells(doc)[1] == ["the single NAT gateway", "no outbound"]
    assert "Evidences: [ICTCLD502 PC 2.2]" in [p.text for p in doc.paragraphs]


def test_response_slot_blank_for_student_modelled_for_assessor():
    doc = Document()
    response_slot(doc, "because one zone is not fault tolerant", "student")
    assert _cells(doc)[0][0] == "[ WRITE YOUR ANSWER HERE ]"
    doc2 = Document()
    response_slot(doc2, "because one zone is not fault tolerant", "assessor")
    assert _cells(doc2)[0][0] == "because one zone is not fault tolerant"


def test_screenshot_slot_describes_for_assessor_prompts_for_student():
    doc = Document()
    screenshot_slot(doc, "the NAT gateway showing State: Available", "student")
    assert _cells(doc)[0][0].startswith("[ PASTE YOUR SCREENSHOT HERE ]")
    doc2 = Document()
    screenshot_slot(doc2, "the NAT gateway showing State: Available", "assessor")
    assert _cells(doc2)[0][0].startswith("SCREENSHOT — ")


# ---------------------------------------------------------------- the scaffolding dials

def test_given_prefills_leading_columns_only():
    """`given` hands over the columns that are context; the rest stay the student's finding."""
    doc = Document()
    design_table(doc, ["Tier", "Meets target?", "Why"],
                 [["Network", "No", "one zone"], ["Compute", "No", "one instance"]],
                 "student", given=1, blank_rows=2)
    assert _cells(doc)[1] == ["Network", "", ""]
    assert _cells(doc)[2] == ["Compute", "", ""]


def test_blank_rows_sets_the_room_offered_not_the_answer_count():
    doc = Document()
    design_table(doc, ["Component"], [["the NAT gateway"]], "student", blank_rows=6)
    assert len(_cells(doc)) == 7                  # header + 6 blank


def test_exemplar_row_renders_worked_and_is_announced():
    doc = Document()
    design_table(doc, ["Component", "Failure mode"],
                 [["an existing subnet", "described, not found"], ["the real answer", "hidden"]],
                 "student", exemplar=1, blank_rows=3)
    assert doc.paragraphs[0].text.startswith("The first row is filled in as an example")
    rows = _cells(doc)
    assert rows[1] == ["an existing subnet", "described, not found"]
    assert rows[2] == ["", ""]                    # the real answer is not handed over
    worked = doc.tables[0].rows[1].cells[0].paragraphs[0].runs[0]
    assert worked.italic is True and str(worked.font.color.rgb) == MODEL


def test_exemplar_is_not_announced_in_the_assessor_copy():
    doc = Document()
    design_table(doc, ["Component"], [["a"], ["b"]], "assessor", exemplar=1)
    assert not any("filled in as an example" in p.text for p in doc.paragraphs)
    assert [r[0] for r in _cells(doc)[1:]] == ["a", "b"]


def test_assessor_table_shows_every_model_row_in_model_colour():
    doc = Document()
    design_table(doc, ["Component"], [["a"], ["b"], ["c"]], "assessor", blank_rows=9)
    assert len(_cells(doc)) == 4                  # header + 3 — blank_rows is a student-copy dial
    run = doc.tables[0].rows[1].cells[0].paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == MODEL
    assert run.font.size == Pt(9)
